from __future__ import annotations

from io import BytesIO
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.schemas.resume import ParsedResumeV2


# =========================================================
# HELPERS
# =========================================================

SKILL_LABELS = {
    "programming_languages": "Languages",
    "frameworks": "Frameworks",
    "libraries": "Libraries",
    "databases": "Databases",
    "cloud": "Cloud",
    "devops": "DevOps",
    "developer_tools": "Developer Tools",
    "ai_ml": "AI / ML",
    "data": "Data",
    "mobile": "Mobile",
    "web": "Web",
    "concepts": "Core CS",
    "soft_skills": "Soft Skills",
    "other": "Other",
}


def _clean(value: str | None) -> str:
    if not value:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _safe_filename(value: str | None) -> str:
    clean = _clean(value) or "Tailored_Resume"
    clean = re.sub(r"[^A-Za-z0-9._ -]+", "", clean)
    clean = re.sub(r"\s+", "_", clean)
    return clean[:100] or "Tailored_Resume"


def _set_cellless_bottom_border(paragraph) -> None:
    p = paragraph._p
    pPr = p.get_or_add_pPr()

    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)


def _set_run_font(run, size: float | None = None, bold: bool | None = None):
    run.font.name = "Arial"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    for style_name in ("Title", "Heading 1", "Heading 2"):
        if style_name in styles:
            styles[style_name].font.name = "Arial"


def _add_header(doc: Document, resume: ParsedResumeV2) -> None:
    info = resume.personal_info

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)

    run = p.add_run(_clean(info.full_name) or "Candidate")
    _set_run_font(run, 15.5, True)

    contact_parts = [
        _clean(info.email),
        _clean(info.phone),
        _clean(info.location),
        _clean(info.linkedin),
        _clean(info.github),
        _clean(info.portfolio),
    ]

    seen = set()
    contact_parts = [
        value for value in contact_parts
        if value and not (value.lower() in seen or seen.add(value.lower()))
    ]

    if contact_parts:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(" | ".join(contact_parts))
        _set_run_font(run, 8.7, False)


def _add_section_heading(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)

    run = p.add_run(title.upper())
    _set_run_font(run, 9.5, True)

    _set_cellless_bottom_border(p)


def _add_bullet(doc: Document, text: str) -> None:
    value = _clean(text)
    if not value:
        return

    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.first_line_indent = Inches(-0.10)
    p.paragraph_format.space_after = Pt(0.6)
    p.paragraph_format.line_spacing = 1.0

    run = p.add_run(value)
    _set_run_font(run, 9.0, False)


def _format_dates(start: str | None, end: str | None, current: bool = False) -> str:
    start_clean = _clean(start)
    end_clean = "Present" if current else _clean(end)

    if start_clean and end_clean:
        return f"{start_clean} - {end_clean}"
    return start_clean or end_clean


def _add_summary(doc: Document, resume: ParsedResumeV2) -> None:
    value = _clean(resume.summary)
    if not value:
        return

    _add_section_heading(doc, "Summary")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(value)
    _set_run_font(run, 9.0, False)


def _add_skills(doc: Document, resume: ParsedResumeV2) -> None:
    rows: list[tuple[str, list[str]]] = []

    for field_name in resume.skills.__class__.model_fields:
        items = list(getattr(resume.skills, field_name, []) or [])
        names = [_clean(getattr(item, "name", "")) for item in items]
        names = [name for name in names if name]

        if names:
            rows.append((
                SKILL_LABELS.get(field_name, field_name.replace("_", " ").title()),
                names,
            ))

    if not rows:
        return

    _add_section_heading(doc, "Skills")

    for label, names in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0.4)

        r1 = p.add_run(f"{label}: ")
        _set_run_font(r1, 8.9, True)

        r2 = p.add_run(", ".join(names))
        _set_run_font(r2, 8.9, False)


def _add_experience(doc: Document, resume: ParsedResumeV2) -> None:
    if not resume.experience:
        return

    _add_section_heading(doc, "Experience")

    for item in resume.experience:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.8)
        p.paragraph_format.space_after = Pt(0.3)

        left = " | ".join(
            value for value in [
                _clean(item.title),
                _clean(item.company),
            ] if value
        )

        right = _format_dates(
            item.start_date,
            item.end_date,
            item.is_current,
        )

        if right:
            text = f"{left}    {right}" if left else right
        else:
            text = left

        run = p.add_run(text)
        _set_run_font(run, 9.2, True)

        meta = " | ".join(
            value for value in [
                _clean(item.location),
                _clean(item.work_mode).title() if _clean(item.work_mode) not in {"", "Unknown"} else "",
                _clean(item.employment_type),
            ] if value
        )

        if meta:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0.2)
            run = p.add_run(meta)
            _set_run_font(run, 8.4, False)

        for bullet in item.bullets:
            _add_bullet(doc, bullet)


def _add_projects(doc: Document, resume: ParsedResumeV2) -> None:
    if not resume.projects:
        return

    _add_section_heading(doc, "Projects")

    for item in resume.projects:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.6)
        p.paragraph_format.space_after = Pt(0.2)

        title = _clean(item.name)
        subtitle = _clean(item.subtitle)

        run = p.add_run(
            f"{title} - {subtitle}" if title and subtitle else (title or subtitle)
        )
        _set_run_font(run, 9.2, True)

        if _clean(item.description):
            _add_bullet(doc, item.description)

        for bullet in item.bullets:
            _add_bullet(doc, bullet)


def _add_education(doc: Document, resume: ParsedResumeV2) -> None:
    if not resume.education:
        return

    _add_section_heading(doc, "Education")

    for item in resume.education:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(0.2)

        institution = _clean(item.institution)
        date_text = _format_dates(item.start_date, item.end_date)

        first = institution
        if date_text:
            first = f"{first}    {date_text}" if first else date_text

        r = p.add_run(first)
        _set_run_font(r, 9.1, True)

        details = []

        degree_field = " in ".join(
            value for value in [
                _clean(item.degree),
                _clean(item.field_of_study),
            ] if value
        )
        if degree_field:
            details.append(degree_field)

        if _clean(item.grade):
            grade_label = _clean(item.grade_type) or "Grade"
            details.append(f"{grade_label}: {_clean(item.grade)}")

        if _clean(item.location):
            details.append(_clean(item.location))

        if details:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0.4)
            r = p.add_run(" | ".join(details))
            _set_run_font(r, 8.8, False)


def _add_certifications(doc: Document, resume: ParsedResumeV2) -> None:
    if not resume.certifications:
        return

    _add_section_heading(doc, "Certifications")

    for item in resume.certifications:
        parts = [
            _clean(item.name),
            _clean(item.issuer),
            _clean(item.issue_date),
        ]
        value = " - ".join(part for part in parts if part)
        if value:
            _add_bullet(doc, value)


def _add_publications_research(doc: Document, resume: ParsedResumeV2) -> None:
    has_content = bool(resume.publications or resume.research)
    if not has_content:
        return

    _add_section_heading(doc, "Publications & Research")

    for item in resume.publications:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.2)
        p.paragraph_format.space_after = Pt(0.2)

        title = _clean(item.title)
        r = p.add_run(title)
        _set_run_font(r, 9.0, True)

        meta = " | ".join(
            value for value in [
                _clean(item.venue),
                str(item.year) if item.year else "",
            ] if value
        )
        if meta:
            p = doc.add_paragraph()
            r = p.add_run(meta)
            _set_run_font(r, 8.6, False)

        if _clean(item.description):
            _add_bullet(doc, item.description)

    for item in resume.research:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.2)
        p.paragraph_format.space_after = Pt(0.2)

        title = _clean(item.title)
        org = _clean(item.organization)

        r = p.add_run(
            f"{title} | {org}" if title and org else (title or org)
        )
        _set_run_font(r, 9.0, True)

        if _clean(item.description):
            _add_bullet(doc, item.description)

        for bullet in item.bullets:
            _add_bullet(doc, bullet)


def _add_achievements(doc: Document, resume: ParsedResumeV2) -> None:
    if not resume.achievements:
        return

    _add_section_heading(doc, "Achievements")

    for item in resume.achievements:
        title = _clean(item.title)
        description = _clean(item.description)

        if title and description:
            _add_bullet(doc, f"{title}: {description}")
        else:
            _add_bullet(doc, title or description)


# =========================================================
# PUBLIC API
# =========================================================

def build_resume_docx(
    resume: ParsedResumeV2,
) -> bytes:
    """
    Build an editable, ATS-friendly DOCX from an already validated
    ParsedResumeV2/tailored resume snapshot.

    The function does not invent or rewrite content.
    """
    doc = Document()
    _configure_document(doc)

    _add_header(doc, resume)
    _add_summary(doc, resume)
    _add_skills(doc, resume)
    _add_experience(doc, resume)
    _add_projects(doc, resume)
    _add_education(doc, resume)
    _add_certifications(doc, resume)
    _add_publications_research(doc, resume)
    _add_achievements(doc, resume)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def resume_docx_filename(
    resume: ParsedResumeV2,
    *,
    company: str | None = None,
    job_title: str | None = None,
) -> str:
    person = _safe_filename(
        resume.personal_info.full_name
    )

    role = _safe_filename(
        job_title
    )

    org = _safe_filename(
        company
    )

    parts = [
        part
        for part in [
            person,
            org if org != "Tailored_Resume" else "",
            role if role != "Tailored_Resume" else "",
        ]
        if part
    ]

    return "_".join(parts) + ".docx"
